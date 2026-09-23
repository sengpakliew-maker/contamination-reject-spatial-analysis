"""Result rendering, plot presentation, and Excel export."""

from __future__ import annotations

import base64
import io
import os
import tempfile

import pandas as pd
import matplotlib.pyplot as plt

def create_device_plot(
    plot_df,
    device_name,
    title_text,
    subtitle_text,
    device_config,
    plot_device_map,
):
    """Render a device map and return its base64 PNG representation."""
    fig, ax = plt.subplots(figsize=(20, 10))
    try:
        plot_device_map(
            df=plot_df,
            ax=ax,
            device_name=device_name,
            device_config=device_config,
            smooth=False,
            sigma=0.6,
        )
        full_title = (
            f"{title_text}\n{subtitle_text}"
            if subtitle_text
            else title_text
        )
        ax.set_title(
            full_title,
            fontsize=16,
            fontweight="bold",
            pad=8,
            linespacing=1.3,
        )
        buffer = io.BytesIO()
        fig.savefig(
            buffer,
            format="png",
            dpi=150,
            bbox_inches="tight",
            facecolor="white",
            pad_inches=0.15,
        )
        buffer.seek(0)
        return base64.b64encode(buffer.read()).decode("utf-8")
    except Exception:
        plt.close(fig)
        raise
    finally:
        if plt.fignum_exists(fig.number):
            plt.close(fig)


def plot_control_from_bytes(ft, image_bytes, width=1100, height=None):
    """Create a Flet image control from PNG bytes."""
    if not image_bytes:
        return None

    image_kwargs = {
        "src": base64.b64encode(image_bytes).decode("utf-8"),
        "width": width,
        "fit": ft.BoxFit.CONTAIN,
    }

    if height is not None:
        image_kwargs["height"] = height

    return ft.Image(**image_kwargs)


def show_method_results(
    ft,
    result_plot_panel,
    method_result,
    plot_control_from_bytes,
):
    """Build and display Method 1 or Method 2 result controls."""
    method = method_result.get("method")
    results = method_result.get("results", [])

    if not results:
        result_plot_panel.content = ft.Text("No method result available.")
        return

    controls = []

    for criteria, devices, result in results:
        device_text = ", ".join(str(device) for device in devices)

        if method == "Method 1":
            controls.append(
                ft.Text(
                    f"Method 1 — DBSCAN — Criteria {criteria} — {device_text}",
                    size=18,
                    weight=ft.FontWeight.BOLD,
                )
            )
            plot_control = plot_control_from_bytes(
                result.get("plot_png"),
                width=1200,
            )
            if plot_control is not None:
                controls.append(plot_control)

            strip_summary = result.get("strip_summary", {})
            cluster_df = result.get("cluster_df")
            cluster_count = len(cluster_df) if cluster_df is not None else 0
            top_strip = result.get("plot_strip_id", "N/A")

            controls.append(
                ft.Text(
                    f"Meaningful clusters={cluster_count:,} | "
                    f"Suspected strips={len(strip_summary):,} | "
                    f"Selected strip={top_strip}"
                )
            )
        else:
            stats = result.get("stats", {})
            controls.append(
                ft.Text(
                    f"Method 2 — Criteria {criteria} — {device_text}",
                    size=18,
                    weight=ft.FontWeight.BOLD,
                )
            )
            plot_control = plot_control_from_bytes(result.get("plot_png"))
            if plot_control is not None:
                controls.append(plot_control)

            map_png = result.get("contribution_map_png")
            if map_png:
                controls.append(
                    ft.Text(
                        "JSD Contribution Map",
                        size=15,
                        weight=ft.FontWeight.BOLD,
                    )
                )
                map_control = plot_control_from_bytes(
                    map_png,
                    width=1500,
                    height=500,
                )
                if map_control is not None:
                    controls.append(map_control)

            controls.append(
                ft.Text(
                    f"Actual JSD: {stats.get('actual_jsd', float('nan')):.6f} | "
                    f"Permutation mean: "
                    f"{stats.get('permutation_mean', float('nan')):.6f} | "
                    f"p%: {stats.get('p_pct', float('nan')):.2f} | "
                    f"{'SIGNIFICANT' if stats.get('significant') else 'NOT SIGNIFICANT'}"
                )
            )
            controls.append(
                ft.Text(
                    "Actual current reject rows traced: "
                    f"{len(result.get('actual_rejects', [])):,}"
                )
            )

        controls.append(ft.Divider())

    result_plot_panel.content = ft.Column(
        controls,
        spacing=10,
        scroll=ft.ScrollMode.AUTO,
    )


def _write_plot_sheet(writer, plot_items):
    """Embed all requested PNG plots in the Plot worksheet."""
    from openpyxl.drawing.image import Image as XLImage
    from openpyxl.styles import Font

    workbook = writer.book
    ws = workbook.create_sheet("Plot")
    ws.column_dimensions["A"].width = 24
    temp_paths = []
    row = 1
    for title, png_bytes in plot_items:
        ws.cell(row=row, column=1, value=title)
        ws.cell(row=row, column=1).font = Font(bold=True, size=14)
        row += 1
        if png_bytes:
            tmp = tempfile.NamedTemporaryFile(delete=False, suffix=".png")
            tmp.write(png_bytes)
            tmp.close()
            temp_paths.append(tmp.name)
            image = XLImage(tmp.name)
            image.width = 900
            image.height = 490
            ws.add_image(image, f"A{row}")
            row += 28
        else:
            ws.cell(row=row, column=1, value="Plot unavailable.")
            row += 2
    writer._method_plot_temp_paths = temp_paths


def _cleanup_plot_temp_paths(writer):
    for path in getattr(writer, "_method_plot_temp_paths", []):
        try:
            os.remove(path)
        except OSError:
            pass


def export_method1_excel(results, output_path):
    """Export DBSCAN Method 1 results and plots to Excel."""
    if isinstance(results, dict):
        results = [("A", results.get("devices", ["Unknown"]), results)]

    cluster_frames = []
    strip_frames = []
    point_frames = []
    plot_items = []

    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        for criteria, devices, result in results:
            device_text = ", ".join(str(d) for d in devices)

            cdf = result.get("cluster_df", pd.DataFrame()).copy()
            sdf = result.get("strip_summary", pd.DataFrame()).copy()
            pdf = result.get("cluster_point_df", pd.DataFrame()).copy()

            cdf.insert(0, "Criteria", criteria) if "Criteria" not in cdf.columns else None
            cdf.insert(1, "Devices", device_text) if "Devices" not in cdf.columns else None
            sdf.insert(0, "Criteria", criteria) if "Criteria" not in sdf.columns else None
            sdf.insert(1, "Devices", device_text) if "Devices" not in sdf.columns else None
            pdf.insert(0, "Criteria", criteria) if "Criteria" not in pdf.columns else None
            pdf.insert(1, "Devices", device_text) if "Devices" not in pdf.columns else None

            cluster_frames.append(cdf)
            strip_frames.append(sdf)
            point_frames.append(pdf)
            plot_items.append((
                f"Method 1 — Criteria {criteria} — {device_text} — DBSCAN",
                result.get("plot_png"),
            ))

        def _write_combined(frames, sheet_name):
            nonempty = [x for x in frames if not x.empty]
            if nonempty:
                pd.concat(nonempty, ignore_index=True).to_excel(writer, sheet_name=sheet_name, index=False)
            else:
                pd.DataFrame().to_excel(writer, sheet_name=sheet_name, index=False)

        _write_combined(cluster_frames, "cluster_df")
        _write_combined(strip_frames, "strip_summary")
        _write_combined(point_frames, "cluster_point_df")
        _write_plot_sheet(writer, plot_items)

    _cleanup_plot_temp_paths(writer)


def export_method2_excel(results, output_path):
    """Export Method 2 results for one or more Criteria groups."""
    if isinstance(results, dict):
        results = [("A", ["Unknown"], results)]
    plot_items = []
    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        stats_frames, bins_frames, reject_frames = [], [], []
        for criteria, devices, result in results:
            device_text = ", ".join(str(d) for d in devices)
            stats_frames.append(pd.DataFrame([result["stats"]]).assign(Criteria=criteria, Devices=device_text))
            bins_frames.append(result["top_bins"].assign(Criteria=criteria, Devices=device_text))
            reject_frames.append(result["actual_rejects"].assign(Criteria=criteria, Devices=device_text))
            plot_items.append((f"Method 2 — Criteria {criteria} — {device_text} — Permutation JSD", result.get("plot_png")))
            plot_items.append((f"Method 2 — Criteria {criteria} — {device_text} — JSD Contribution", result.get("contribution_map_png")))
        pd.concat(stats_frames, ignore_index=True).to_excel(writer, sheet_name="Summary", index=False)
        pd.concat(bins_frames, ignore_index=True).to_excel(writer, sheet_name="Top Bins", index=False)
        pd.concat(reject_frames, ignore_index=True).to_excel(writer, sheet_name="Actual Rejects", index=False)
        _write_plot_sheet(writer, plot_items)
    _cleanup_plot_temp_paths(writer)
